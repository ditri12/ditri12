#!/usr/bin/env python3
"""Build the laboratory-supplier database.

Population: every supplier/brand listed on the catalog pages of Dutscher (dutscher.com/brands),
Thomas Scientific (thomassci.com/browse-by/suppliers) and Quartzy (quartzy.com/suppliers,
info.quartzy.com/new-suppliers) - research/catalog_suppliers.json (built by crawl/parse_lists.py).

Enrichment layers (all optional, research/enrich/*.json):
  prior_research.json   verified research carried over from the earlier top-100 database
  resolved.json         Wikidata / Clearbit resolution (website, country, HQ, parent, employees, revenue)
  catalog_products.json products shown for the supplier in the catalogs
  israel_matches.json   brand mentions found on Israeli distributor sites (crawl/match_israel.py)
  site_contacts.json    phone / e-mail found on the supplier's own site
  research/israel/distributors.json  registry of Israeli distributors (contacts)
  research/overrides.json            hand-curated corrections, keyed by master key

Outputs: data/suppliers.{json,csv,xlsx,sql,sqlite} and index.html (viewer).
"""
import csv
import datetime as dt
import json
import re
import sqlite3
from pathlib import Path

ROOT = Path(__file__).resolve().parent
RESEARCH = ROOT / "research"
ENRICH = RESEARCH / "enrich"
DATA = ROOT / "data"
TEMPLATE = ROOT / "template" / "index.template.html"
OUT_HTML = ROOT / "index.html"

CATALOGS = {
    "dutscher": {"name": "Dutscher", "name_he": "Dutscher (צרפת)", "url": "https://www.dutscher.com/brands", "short": "D"},
    "thomassci": {"name": "Thomas Scientific", "name_he": "Thomas Scientific (ארה\"ב)", "url": "https://www.thomassci.com/browse-by/suppliers", "short": "T"},
    "quartzy": {"name": "Quartzy", "name_he": "Quartzy (ארה\"ב)", "url": "https://www.quartzy.com/suppliers", "short": "Q"},
    "quartzy_new": {"name": "Quartzy – new suppliers", "name_he": "Quartzy – ספקים חדשים", "url": "https://info.quartzy.com/new-suppliers", "short": "Q+"},
}
PRESENCE = ["subsidiary", "distributor", "parent_subsidiary", "none_found", "unclear", "not_found_in_sweep", "pending"]
PRESENCE_HE = {
    "subsidiary": "חברה-בת בישראל",
    "distributor": "מפיץ/נציג בישראל",
    "parent_subsidiary": "דרך חברת-האם בישראל",
    "none_found": "לא נמצא ערוץ ישראלי",
    "unclear": "לא ברור",
    "not_found_in_sweep": "לא נמצא בסריקת המפיצים",
    "pending": "טרם נבדק",
}
CHECK_HE = {"verified": "אומת", "partial": "חיפוש חלקי", "sweep": "סריקת אתרי מפיצים", "unverified": "לא אומת", "pending": "טרם נבדק"}
ETYPE_HE = {"subsidiary": "חברה-בת", "distributor": "מפיץ", "service": "שירות", "office": "משרד"}
ESTAT_HE = {"current": "עדכני", "older_evidence": "ראיה ישנה", "no_longer": "לא עוד", "new": "חדש"}
SIZE_SRC_HE = {"prior": "מחקר קודם מאומת", "wikidata": "Wikidata", "site": "אתר החברה", "manual": "ידני", "estimate": "הערכה"}
COUNTRY_HE = {
    "United States of America": "ארה\"ב", "United States": "ארה\"ב", "USA": "ארה\"ב", "Germany": "גרמניה", "United Kingdom": "בריטניה",
    "France": "צרפת", "Switzerland": "שווייץ", "Japan": "יפן", "China": "סין", "People's Republic of China": "סין", "Italy": "איטליה",
    "Netherlands": "הולנד", "Kingdom of the Netherlands": "הולנד", "Denmark": "דנמרק", "Sweden": "שוודיה", "Austria": "אוסטריה",
    "Belgium": "בלגיה", "Spain": "ספרד", "Canada": "קנדה", "Israel": "ישראל", "India": "הודו", "South Korea": "דרום קוריאה",
    "Taiwan": "טייוואן", "Australia": "אוסטרליה", "Czech Republic": "צ'כיה", "Czechia": "צ'כיה", "Poland": "פולין", "Finland": "פינלנד",
    "Norway": "נורווגיה", "Ireland": "אירלנד", "Republic of Ireland": "אירלנד", "Singapore": "סינגפור", "Hungary": "הונגריה", "Portugal": "פורטוגל",
    "Lithuania": "ליטא", "Luxembourg": "לוקסמבורג", "New Zealand": "ניו זילנד", "Brazil": "ברזיל", "Mexico": "מקסיקו", "Turkey": "טורקיה",
    "Hong Kong": "הונג קונג", "Malaysia": "מלזיה", "Thailand": "תאילנד", "Slovakia": "סלובקיה", "Slovenia": "סלובניה", "Greece": "יוון",
    "Russia": "רוסיה", "South Africa": "דרום אפריקה", "Estonia": "אסטוניה", "Latvia": "לטביה", "Iceland": "איסלנד", "Liechtenstein": "ליכטנשטיין",
}


CURRENCY_CODE = {"United States dollar": "USD", "euro": "EUR", "pound sterling": "GBP", "Swiss franc": "CHF", "Japanese yen": "JPY",
                 "Chinese yuan": "CNY", "renminbi": "CNY", "Danish krone": "DKK", "Swedish krona": "SEK", "Canadian dollar": "CAD",
                 "Indian rupee": "INR", "South Korean won": "KRW", "New Taiwan dollar": "TWD", "Czech koruna": "CZK", "Polish złoty": "PLN", "Polish zloty": "PLN",
                 "Norwegian krone": "NOK", "Australian dollar": "AUD", "Israeli new shekel": "ILS", "Hong Kong dollar": "HKD", "Singapore dollar": "SGD",
                 "Brazilian real": "BRL", "Mexican peso": "MXN", "Turkish lira": "TRY", "South African rand": "ZAR", "Hungarian forint": "HUF"}
FX_TO_USD = {"USD": 1.0, "EUR": 1.08, "GBP": 1.27, "CHF": 1.12, "JPY": 0.0067, "CNY": 0.14, "DKK": 0.145, "SEK": 0.095, "CAD": 0.73, "INR": 0.012,
             "KRW": 0.00072, "TWD": 0.031, "CZK": 0.043, "PLN": 0.25, "NOK": 0.093, "AUD": 0.65, "ILS": 0.27, "HKD": 0.128, "SGD": 0.74, "BRL": 0.18,
             "MXN": 0.055, "TRY": 0.03, "ZAR": 0.055, "HUF": 0.0027}


def load_json(path, default):
    if Path(path).exists():
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return default


def domain_of(url):
    if not url:
        return ""
    m = re.match(r"https?://([^/]+)", url)
    return (m.group(1) if m else url).lower().replace("www.", "")


def clean_phone(p):
    return re.sub(r"\s+", " ", (p or "")).strip()


def first(*vals):
    for v in vals:
        if v not in (None, "", [], {}):
            return v
    return None


def build_row(m, prior, res, il_matches, il_reg, prods, site_contact, ov):
    key = m["key"]
    prior = prior or {}
    res = res or {}
    wd = res.get("wikidata") or {}
    cb = res.get("clearbit") or {}
    ov = ov or {}
    r = {"id": m["id"], "key": key, "name": ov.get("name") or m["name"], "names": m.get("names") or [m["name"]]}
    # catalogs
    cats = {}
    for c, src in (m.get("sources") or {}).items():
        cats[c] = {"url": src.get("url"), "own_brand": bool(src.get("own_brand")), "description": src.get("description"),
                   "catalog_link": src.get("catalog_link")}
    r["catalogs"] = cats
    r["n_catalogs"] = len(cats)
    r["own_brand_of"] = "Dutscher" if any(c.get("own_brand") for c in cats.values()) else None
    # website
    web, web_src = None, None
    if ov.get("website"):
        web, web_src = ov["website"], "manual"
    elif prior.get("website"):
        web, web_src = prior["website"], "prior"
    elif wd.get("website"):
        web, web_src = wd["website"], "wikidata"
    elif cb.get("domain") and (cb.get("score") or 0) >= 0.8:
        web, web_src = "https://" + cb["domain"], "clearbit"
    r["website"] = web
    r["website_source"] = web_src
    r["website_candidate"] = ("https://" + cb["domain"]) if (not web and cb.get("domain")) else None
    r["website_candidate_score"] = cb.get("score") if r["website_candidate"] else None
    # identity
    country = first(ov.get("country"), prior.get("country"), (wd.get("country") or [None])[0])
    r["country"] = country
    r["country_he"] = COUNTRY_HE.get(country or "", country or "")
    r["hq"] = first(ov.get("hq"), prior.get("hq"), (wd.get("hq") or [None])[0])
    r["parent"] = first(ov.get("parent"), prior.get("parent"), (wd.get("parent") or [None])[0])
    r["inception"] = first(ov.get("inception"), prior.get("founded"), wd.get("inception"))
    r["aka"] = prior.get("aka")
    r["ownership"] = first(ov.get("ownership"), prior.get("ownership"))
    r["industry"] = first(ov.get("industry"), prior.get("category"), ", ".join(wd.get("industry") or []) or None)
    desc = first(ov.get("description"), (cats.get("quartzy_new") or {}).get("description"), wd.get("description"))
    r["description"] = desc
    r["wikidata"] = wd.get("qid")
    r["wikidata_score"] = wd.get("score")
    # size
    emp = None
    if ov.get("employees") is not None:
        emp = {"value": ov["employees"], "year": ov.get("employees_year"), "scope": ov.get("employees_scope", "company"), "source": "manual", "note": ov.get("employees_note")}
    elif prior.get("employees") is not None:
        emp = {"value": prior["employees"], "year": prior.get("employees_year"), "scope": prior.get("employees_scope") or "company", "source": "prior", "note": prior.get("employees_note")}
    elif wd.get("employees") and wd["employees"].get("value"):
        emp = {"value": int(wd["employees"]["value"]), "year": wd["employees"].get("year"), "scope": "company", "source": "wikidata", "note": None}
    rev = None
    if ov.get("revenue_usd_m") is not None:
        rev = {"usd_m": ov["revenue_usd_m"], "original": ov.get("revenue_original"), "year": ov.get("revenue_year"), "scope": ov.get("revenue_scope", "company"), "source": "manual", "note": ov.get("revenue_note")}
    elif prior.get("revenue_usd_m") is not None:
        rev = {"usd_m": prior["revenue_usd_m"], "original": prior.get("revenue_original"), "year": prior.get("revenue_year"), "scope": prior.get("revenue_scope") or "company", "source": "prior", "note": prior.get("revenue_note")}
    elif wd.get("revenue") and wd["revenue"].get("value"):
        v = wd["revenue"]
        cur = CURRENCY_CODE.get(v.get("currency") or "", v.get("currency") or "")
        fx = FX_TO_USD.get(cur)
        usd_m = round(v["value"] * fx / 1e6, 1) if fx else None
        rev = {"usd_m": usd_m, "original": f"{cur} {v['value'] / 1e6:,.1f} m" if cur else f"{v['value'] / 1e6:,.1f} m",
               "year": v.get("year"), "scope": "company", "source": "wikidata", "note": "Wikidata (P2139); conversion to USD at approximate 2026 rates" if fx and cur != "USD" else "Wikidata (P2139)"}
    r["employees"] = emp
    r["revenue"] = rev
    r["size_basis"] = first(ov.get("size_basis"), prior.get("size_basis"), "reported" if (emp or rev) else "unknown")
    # products
    plist = []
    for p in (ov.get("products") or []):
        plist.append({"name": p if isinstance(p, str) else p.get("name"), "source": "manual"})
    if not plist:
        for p in (prior.get("products") or []):
            plist.append({"name": p if isinstance(p, str) else p.get("name"), "desc": None if isinstance(p, str) else p.get("desc"), "source": "prior"})
    if not plist and prods:
        for p in prods:
            plist.append({"name": p.get("name"), "source": p.get("source", "catalog"), "url": p.get("url"), "catalog": p.get("catalog")})
    r["products"] = [p for p in plist if p.get("name")][:5]
    # Israel
    ents = []
    for e in (ov.get("israel_entities") or prior.get("israel_entities") or []):
        e = dict(e)
        e.setdefault("match_basis", "prior_research")
        ents.append(e)
    seen_hosts = {domain_of(e.get("website")) for e in ents if e.get("website")}
    seen_names = {re.sub(r"[^a-z0-9]", "", (e.get("name") or "").lower())[:10] for e in ents}
    sweep_hits = []
    for h in (il_matches or []):
        reg = il_reg.get(h.get("distributor")) if h.get("distributor") else None
        reg = reg or {}
        host = domain_of(reg.get("website") or h.get("website") or h.get("site"))
        if host and host in seen_hosts:
            continue
        nm = h.get("name") or reg.get("name") or h.get("site")
        nkey = re.sub(r"[^a-z0-9]", "", (nm or "").lower())[:10]
        if nkey in seen_names:
            continue
        if host:
            seen_hosts.add(host)
        seen_names.add(nkey)
        ents.append({"name": nm, "name_he": h.get("name_he") or reg.get("name_he"), "type": h.get("type") or "distributor",
                     "scope": h.get("scope") or "listed as a brand on the distributor's site",
                     "website": h.get("website") or reg.get("website") or ("https://" + host if host else None),
                     "phone": h.get("phone") or reg.get("phone"), "email": h.get("email") or reg.get("email"),
                     "address": h.get("address") or reg.get("address"), "contact_url": h.get("contact_url") or reg.get("contact_url"),
                     "source": h.get("source") or h.get("page_url"),
                     "evidence_year": h.get("evidence_year") if "evidence_year" in h else 2026, "status": h.get("status") or "current",
                     "match_basis": h.get("match_basis") or "distributor_site_sweep",
                     "match_confidence": h.get("match_confidence") or h.get("confidence"), "match_text": h.get("evidence")})
        sweep_hits.append(h)
    if ov.get("israel_presence"):
        presence = ov["israel_presence"]
        check = ov.get("israel_check", "verified")
    elif prior.get("israel_presence"):
        presence = prior["israel_presence"]
        check = prior.get("israel_check") or "verified"
        if presence in ("none_found", "unclear") and sweep_hits:
            presence = "distributor"
            check = "partial"
    elif sweep_hits:
        presence = "distributor"
        check = "sweep"
    else:
        presence = "not_found_in_sweep"
        check = "sweep"
    r["israel_presence"] = presence
    r["israel_presence_he"] = PRESENCE_HE[presence]
    r["israel_check"] = check
    r["israel_check_he"] = CHECK_HE[check]
    r["israel_entities"] = ents
    yrs = [e.get("evidence_year") for e in ents if e.get("evidence_year")]
    r["israel_evidence_year"] = max(yrs) if yrs else None
    r["israel_note"] = first(ov.get("israel_note"), prior.get("israel_note"))
    r["israel_current"] = first(ov.get("israel_current"), prior.get("israel_current"), "confirmed_current" if sweep_hits else None)
    # contacts
    c = {}
    pc = prior.get("contact") or {}
    sc = site_contact or {}
    oc = ov.get("contact") or {}
    c["phone"] = clean_phone(first(oc.get("phone"), pc.get("phone"), sc.get("phone")))
    c["email"] = first(oc.get("email"), pc.get("email"), sc.get("email"))
    c["address"] = first(oc.get("address"), pc.get("address"), sc.get("address"), r["hq"])
    c["contact_url"] = first(oc.get("contact_url"), pc.get("contact_url"), sc.get("contact_url"))
    c["source"] = first(oc.get("source"), pc.get("source"), sc.get("source"))
    c["note"] = first(oc.get("note"), pc.get("note"), sc.get("note"))
    r["contact"] = c
    # sources
    srcs = []
    for cat, cinfo in cats.items():
        if cinfo.get("url"):
            srcs.append(cinfo["url"])
    srcs += [s for s in (prior.get("sources") or []) if isinstance(s, str)]
    if wd.get("qid"):
        srcs.append("https://www.wikidata.org/wiki/" + wd["qid"])
    srcs += [h.get("page_url") for h in sweep_hits if h.get("page_url")]
    if sc.get("source"):
        srcs.append(sc["source"])
    srcs += [h.get("source") for h in sweep_hits if h.get("source")]
    r["sources"] = list(dict.fromkeys(s for s in srcs if s))
    r["prior_verified"] = bool(prior)
    r["prior_verified_at"] = prior.get("verified_at")
    return r


def load_rows():
    master = load_json(RESEARCH / "catalog_suppliers.json", [])
    prior = load_json(ENRICH / "prior_research.json", {})
    resolved = {x["key"]: x for x in load_json(ENRICH / "resolved.json", [])}
    il = load_json(ENRICH / "israel_matches.json", {})
    for k, hits in load_json(ENRICH / "israel_knowledge.json", {}).items():
        il.setdefault(k, [])
        il[k] = list(hits) + il[k]
    il_reg = load_json(RESEARCH / "israel" / "distributors.json", {})
    prods = load_json(ENRICH / "catalog_products.json", {})
    contacts = load_json(ENRICH / "site_contacts.json", {})
    overrides = load_json(RESEARCH / "overrides.json", {})
    rows = []
    for m in master:
        k = m["key"]
        rows.append(build_row(m, prior.get(k), resolved.get(k), il.get(k), il_reg, prods.get(k), contacts.get(k), overrides.get(k)))
    return rows


# ---------------------------------------------------------------- outputs
def entity_key(name):
    n = re.sub(r"\(.*?\)", " ", (name or "").lower())
    toks = [t for t in re.split(r"[^a-z0-9&]+", n) if t]
    skip = {"dr", "dr.", "the", "y.a.", "s.l.", "uri", "im", "i.m."}
    for t in toks:
        if t in skip or (len(t) <= 2 and "." in t):
            continue
        return t
    return n.strip()


def distributor_index(rows):
    idx = {}
    for r in rows:
        for e in r["israel_entities"]:
            k = domain_of(e.get("website")) or entity_key(e.get("name"))
            d = idx.setdefault(k, {"key": k, "name": e.get("name"), "name_he": e.get("name_he"), "type": e.get("type"), "website": e.get("website"),
                                   "phone": e.get("phone"), "email": e.get("email"), "address": e.get("address"), "contact_url": e.get("contact_url"),
                                   "suppliers": []})
            for f in ("phone", "email", "address", "contact_url", "name_he", "website"):
                if not d.get(f) and e.get(f):
                    d[f] = e[f]
            d["suppliers"].append({"id": r["id"], "name": r["name"], "status": e.get("status"), "evidence_year": e.get("evidence_year"), "match_basis": e.get("match_basis")})
    out = sorted(idx.values(), key=lambda d: -len(d["suppliers"]))
    for d in out:
        d["count"] = len(d["suppliers"])
    return out


def write_json(rows, generated):
    DATA.mkdir(exist_ok=True)
    payload = {
        "title": "Laboratory suppliers listed by Dutscher, Thomas Scientific and Quartzy",
        "title_he": "ספקי ציוד מעבדה המופיעים בקטלוגים של Dutscher, Thomas Scientific ו-Quartzy",
        "generated": generated, "count": len(rows), "catalogs": CATALOGS,
        "israel_presence_values": PRESENCE_HE, "israel_check_values": CHECK_HE,
        "suppliers": rows, "israeli_entities_index": distributor_index(rows),
    }
    with open(DATA / "suppliers.json", "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    return payload


CSV_COLUMNS = ["id", "name", "catalogs", "dutscher_url", "thomassci_url", "quartzy", "quartzy_new", "own_brand_of", "website", "website_source",
               "country", "country_he", "hq", "parent", "inception", "industry", "description",
               "employees", "employees_year", "employees_scope", "employees_source", "revenue_usd_m", "revenue_original", "revenue_year", "revenue_scope", "revenue_source",
               "size_basis", "product_1", "product_2", "product_3", "products_source",
               "israel_presence", "israel_presence_he", "israel_check", "israel_check_he", "israel_evidence_year", "israel_entities", "israel_entities_contacts", "israel_note",
               "contact_phone", "contact_email", "contact_address", "contact_url", "contact_source", "wikidata", "sources"]


def flat_row(r):
    e = r.get("employees") or {}
    v = r.get("revenue") or {}
    prods = r.get("products") or []
    ents = r.get("israel_entities") or []
    cat_short = "+".join(CATALOGS[c]["short"] for c in CATALOGS if c in r["catalogs"])
    return {
        "id": r["id"], "name": r["name"], "catalogs": cat_short,
        "dutscher_url": (r["catalogs"].get("dutscher") or {}).get("url"), "thomassci_url": (r["catalogs"].get("thomassci") or {}).get("url"),
        "quartzy": "yes" if "quartzy" in r["catalogs"] else "", "quartzy_new": "yes" if "quartzy_new" in r["catalogs"] else "",
        "own_brand_of": r.get("own_brand_of"), "website": r.get("website"), "website_source": r.get("website_source"),
        "country": r.get("country"), "country_he": r.get("country_he"), "hq": r.get("hq"), "parent": r.get("parent"), "inception": r.get("inception"),
        "industry": r.get("industry"), "description": r.get("description"),
        "employees": e.get("value"), "employees_year": e.get("year"), "employees_scope": e.get("scope"), "employees_source": e.get("source"),
        "revenue_usd_m": v.get("usd_m"), "revenue_original": v.get("original"), "revenue_year": v.get("year"), "revenue_scope": v.get("scope"), "revenue_source": v.get("source"),
        "size_basis": r.get("size_basis"),
        "product_1": prods[0]["name"] if len(prods) > 0 else "", "product_2": prods[1]["name"] if len(prods) > 1 else "", "product_3": prods[2]["name"] if len(prods) > 2 else "",
        "products_source": prods[0].get("source") if prods else "",
        "israel_presence": r["israel_presence"], "israel_presence_he": r["israel_presence_he"], "israel_check": r["israel_check"], "israel_check_he": r["israel_check_he"],
        "israel_evidence_year": r.get("israel_evidence_year"),
        "israel_entities": "; ".join(x.get("name") or "" for x in ents),
        "israel_entities_contacts": "; ".join(f"{x.get('name')}: {x.get('phone') or ''} {x.get('email') or ''}".strip() for x in ents),
        "israel_note": r.get("israel_note"),
        "contact_phone": r["contact"].get("phone"), "contact_email": r["contact"].get("email"), "contact_address": r["contact"].get("address"),
        "contact_url": r["contact"].get("contact_url"), "contact_source": r["contact"].get("source"), "wikidata": r.get("wikidata"),
        "sources": " | ".join(r.get("sources") or []),
    }


def write_csv(rows):
    with open(DATA / "suppliers.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow(flat_row(r))


def write_xlsx(rows, generated):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E78")

    def sheet(ws, headers, widths):
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=i)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    ws = wb.active
    ws.title = "Suppliers"
    headers = ["#", "ספק/יצרן", "קטלוגים", "Dutscher", "Thomas Sci.", "Quartzy", "מותג-בית של", "אתר", "מקור האתר", "מדינה", "מטה", "חברת-אם", "שנת ייסוד", "תחום",
               "תיאור", "עובדים", "שנת נתון", "היקף", "מקור", "הכנסות (מיליוני $)", "הכנסות (מקור)", "שנה", "היקף", "מקור", "בסיס הגודל",
               "מוצר 1", "מוצר 2", "מוצר 3", "מקור המוצרים", "ייצוג בישראל", "סטטוס בדיקה", "שנת ראיה", "נציגים בישראל", "אנשי קשר – נציגים", "הערה (ישראל)",
               "טלפון", "דוא\"ל", "כתובת", "דף יצירת קשר", "מקור אנשי קשר", "Wikidata", "מקורות"]
    widths = [5, 30, 8, 12, 12, 8, 12, 30, 9, 12, 18, 20, 8, 20, 50, 10, 8, 9, 9, 14, 7, 7, 9, 9, 10, 28, 28, 28, 10, 20, 14, 8, 40, 50, 40, 16, 24, 36, 30, 20, 10, 60]
    sheet(ws, headers, widths)
    for r in rows:
        fr = flat_row(r)
        ws.append([fr["id"], fr["name"], fr["catalogs"], fr["dutscher_url"], fr["thomassci_url"], fr["quartzy"] + (" +new" if fr["quartzy_new"] else ""), fr["own_brand_of"],
                   fr["website"], fr["website_source"], fr["country_he"], fr["hq"], fr["parent"], fr["inception"], fr["industry"], fr["description"],
                   fr["employees"], fr["employees_year"], fr["employees_scope"], fr["employees_source"], fr["revenue_usd_m"], fr["revenue_original"], fr["revenue_year"], fr["revenue_scope"], fr["revenue_source"], fr["size_basis"],
                   fr["product_1"], fr["product_2"], fr["product_3"], fr["products_source"], fr["israel_presence_he"], fr["israel_check_he"], fr["israel_evidence_year"],
                   fr["israel_entities"], fr["israel_entities_contacts"], fr["israel_note"], fr["contact_phone"], fr["contact_email"], fr["contact_address"], fr["contact_url"], fr["contact_source"], fr["wikidata"], fr["sources"]])
    n = len(rows)
    # Israel entities sheet
    ws2 = wb.create_sheet("Israel entities")
    sheet(ws2, ["# ספק", "ספק/יצרן", "נציג בישראל", "שם בעברית", "סוג", "היקף", "אתר", "טלפון", "דוא\"ל", "כתובת", "דף יצירת קשר", "שנת ראיה", "סטטוס", "בסיס ההתאמה", "ביטחון", "מקור"],
          [7, 28, 34, 22, 10, 36, 28, 16, 26, 36, 30, 8, 10, 18, 9, 60])
    for r in rows:
        for e in r["israel_entities"]:
            ws2.append([r["id"], r["name"], e.get("name"), e.get("name_he"), ETYPE_HE.get(e.get("type"), e.get("type")), e.get("scope"), e.get("website"),
                        e.get("phone"), e.get("email"), e.get("address"), e.get("contact_url"), e.get("evidence_year"), ESTAT_HE.get(e.get("status"), e.get("status")),
                        e.get("match_basis"), e.get("match_confidence"), e.get("source")])
    # Products sheet
    ws3 = wb.create_sheet("Products")
    sheet(ws3, ["# ספק", "ספק/יצרן", "מוצר", "תיאור", "מקור", "קישור"], [7, 28, 50, 50, 10, 60])
    for r in rows:
        for p in r["products"]:
            ws3.append([r["id"], r["name"], p.get("name"), p.get("desc"), p.get("source"), p.get("url")])
    # By Israeli entity
    ws4 = wb.create_sheet("By Israeli entity")
    sheet(ws4, ["נציג בישראל", "שם בעברית", "סוג", "אתר", "טלפון", "דוא\"ל", "כתובת", "מספר ספקים", "ספקים"], [34, 22, 10, 28, 16, 26, 36, 10, 90])
    for d in distributor_index(rows):
        ws4.append([d["name"], d.get("name_he"), ETYPE_HE.get(d.get("type"), d.get("type")), d.get("website"), d.get("phone"), d.get("email"), d.get("address"), d["count"],
                    "; ".join(s["name"] for s in d["suppliers"])])
    # Summary with formulas
    ws5 = wb.create_sheet("Summary")
    ws5.sheet_view.rightToLeft = True
    ws5.column_dimensions["A"].width = 42
    ws5.column_dimensions["B"].width = 14
    ws5.append(["סיכום", generated])
    ws5["A1"].font = Font(bold=True, size=13)
    ws5.append(["מספר ספקים", f"=COUNTA(Suppliers!B2:B{n + 1})"])
    for c in CATALOGS:
        ws5.append([f"מופיעים ב-{CATALOGS[c]['name']}", f'=COUNTIF(Suppliers!{"D" if c == "dutscher" else "E" if c == "thomassci" else "F"}2:{"D" if c == "dutscher" else "E" if c == "thomassci" else "F"}{n + 1},"{"*new*" if c == "quartzy_new" else "?*"}")'])
    for p in PRESENCE:
        ws5.append([PRESENCE_HE[p], f'=COUNTIF(Suppliers!AD2:AD{n + 1},"{PRESENCE_HE[p]}")'])
    ws5.append(["עם אתר מזוהה", f'=COUNTIF(Suppliers!H2:H{n + 1},"http*")'])
    ws5.append(["עם נתון עובדים", f"=COUNT(Suppliers!P2:P{n + 1})"])
    ws5.append(["עם נתון הכנסות", f"=COUNT(Suppliers!T2:T{n + 1})"])
    ws5.append(["עם 3 מוצרים", f'=COUNTIF(Suppliers!AB2:AB{n + 1},"?*")'])
    ws5.append(["עם טלפון", f'=COUNTIF(Suppliers!AJ2:AJ{n + 1},"?*")'])
    ws5.append(["חציון עובדים", f"=MEDIAN(Suppliers!P2:P{n + 1})"])
    for row in ws5.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
    wb.save(DATA / "suppliers.xlsx")


def sql_quote(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


SCHEMA = """\
CREATE TABLE suppliers (
  id INTEGER PRIMARY KEY, key TEXT, name TEXT, catalogs TEXT, dutscher_url TEXT, thomassci_url TEXT, quartzy TEXT, quartzy_new TEXT, own_brand_of TEXT,
  website TEXT, website_source TEXT, country TEXT, country_he TEXT, hq TEXT, parent TEXT, inception TEXT, industry TEXT, description TEXT,
  employees INTEGER, employees_year TEXT, employees_scope TEXT, employees_source TEXT,
  revenue_usd_m REAL, revenue_original TEXT, revenue_year TEXT, revenue_scope TEXT, revenue_source TEXT, size_basis TEXT,
  israel_presence TEXT, israel_presence_he TEXT, israel_check TEXT, israel_check_he TEXT, israel_evidence_year INTEGER, israel_note TEXT,
  contact_phone TEXT, contact_email TEXT, contact_address TEXT, contact_url TEXT, contact_source TEXT, wikidata TEXT
);
CREATE TABLE products (supplier_id INTEGER REFERENCES suppliers(id), position INTEGER, name TEXT, description TEXT, source TEXT, url TEXT);
CREATE TABLE israel_entities (supplier_id INTEGER REFERENCES suppliers(id), name TEXT, name_he TEXT, type TEXT, scope TEXT, website TEXT, phone TEXT, email TEXT, address TEXT, contact_url TEXT, evidence_year INTEGER, status TEXT, match_basis TEXT, match_confidence TEXT, source TEXT);
CREATE TABLE sources (supplier_id INTEGER REFERENCES suppliers(id), url TEXT);
CREATE INDEX idx_suppliers_country ON suppliers(country);
CREATE INDEX idx_suppliers_presence ON suppliers(israel_presence);
CREATE INDEX idx_entities_name ON israel_entities(name);
"""


def write_sql_and_sqlite(rows, generated):
    stmts = []
    for r in rows:
        fr = flat_row(r)
        cols = ["id", "key", "name", "catalogs", "dutscher_url", "thomassci_url", "quartzy", "quartzy_new", "own_brand_of", "website", "website_source", "country", "country_he", "hq", "parent", "inception", "industry", "description",
                "employees", "employees_year", "employees_scope", "employees_source", "revenue_usd_m", "revenue_original", "revenue_year", "revenue_scope", "revenue_source", "size_basis",
                "israel_presence", "israel_presence_he", "israel_check", "israel_check_he", "israel_evidence_year", "israel_note",
                "contact_phone", "contact_email", "contact_address", "contact_url", "contact_source", "wikidata"]
        vals = [r["id"], r["key"]] + [fr[c] for c in cols[2:]]
        stmts.append(f"INSERT INTO suppliers ({', '.join(cols)}) VALUES ({', '.join(sql_quote(v) for v in vals)});")
        for i, p in enumerate(r["products"], 1):
            stmts.append(f"INSERT INTO products VALUES ({r['id']}, {i}, {sql_quote(p.get('name'))}, {sql_quote(p.get('desc'))}, {sql_quote(p.get('source'))}, {sql_quote(p.get('url'))});")
        for e in r["israel_entities"]:
            stmts.append("INSERT INTO israel_entities VALUES (" + ", ".join(sql_quote(v) for v in [r["id"], e.get("name"), e.get("name_he"), e.get("type"), e.get("scope"), e.get("website"), e.get("phone"), e.get("email"), e.get("address"), e.get("contact_url"), e.get("evidence_year"), e.get("status"), e.get("match_basis"), e.get("match_confidence"), e.get("source")]) + ");")
        for s in r["sources"]:
            stmts.append(f"INSERT INTO sources VALUES ({r['id']}, {sql_quote(s)});")
    sql = f"-- Laboratory suppliers database, generated {generated}\nBEGIN;\n{SCHEMA}\n" + "\n".join(stmts) + "\nCOMMIT;\n"
    with open(DATA / "suppliers.sql", "w", encoding="utf-8") as f:
        f.write(sql)
    db = DATA / "suppliers.sqlite"
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    con.executescript(sql)
    con.close()


def write_html(payload):
    if not TEMPLATE.exists():
        return
    tpl = TEMPLATE.read_text(encoding="utf-8")
    data = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    OUT_HTML.write_text(tpl.replace("__DATA_JSON__", data), encoding="utf-8")


def main():
    generated = dt.date.today().isoformat()
    rows = load_rows()
    payload = write_json(rows, generated)
    write_csv(rows)
    write_xlsx(rows, generated)
    write_sql_and_sqlite(rows, generated)
    write_html(payload)
    from collections import Counter
    print(f"built {len(rows)} suppliers")
    print("presence:", dict(Counter(r["israel_presence"] for r in rows)))
    print("website:", sum(1 for r in rows if r["website"]), "employees:", sum(1 for r in rows if r["employees"]), "revenue:", sum(1 for r in rows if r["revenue"]),
          "products:", sum(1 for r in rows if r["products"]), "phone:", sum(1 for r in rows if r["contact"].get("phone")))


if __name__ == "__main__":
    main()
